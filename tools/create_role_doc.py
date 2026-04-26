# -*- coding: utf-8 -*-
"""Erstellt eine Word-Datei mit allen Grimmhain-Rollen, Tier und Kurzbeschreibungen."""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

# ─── DATEN ────────────────────────────────────────────────────────────────────

ROLES = [
    # (DE-Name, EN-Name, Tier, DE-Beschreibung ≤20 Wörter, EN-Beschreibung ≤20 Wörter)
    ("Loki", "Loki", "0.1",
     "Verbindet einmalig zwei Spieler als Liebende oder verflucht zwei als ewige Rivalen.",
     "Once: binds two players as lovers or curses two as eternal rivals."),

    ("Nachtwächter", "Night Warden", "X",
     "Spürt wenn ein Nachbar kein Dorfbewohner ist. Alarm ertönt öffentlich.",
     "Senses when a neighbor doesn't belong. Alarm bells ring publicly."),

    ("Die Gebundenen", "The Bound", "0.15",
     "Lernen in Nacht 1 alle anderen Gebundenen kennen.",
     "Wake in night 1 and learn all other Bound players."),

    ("Waldhexe", "Witch of the Woods", "0.9",
     "Sieht das Schicksal des Opfers. Einmalig retten oder einen anderen verurteilen.",
     "Sees the victim's fate. Once per game: spare them or doom another."),

    ("Rattenfänger", "Pied Piper", "1.1",
     "Verzaubert jede Nacht 1–2 Spieler. Gewinnt wenn alle lebenden Spieler verzaubert sind.",
     "Charms 1–2 players each night. Wins when all living players are charmed."),

    ("Sensenträger", "Reaper", "X",
     "Beim Tod: erntet eine letzte Seele ihrer Wahl.",
     "Upon death: harvests one final soul of her choice."),

    ("Wolfskind", "Wolf Child", "0.3",
     "Wählt einen Feind. Stirbt dieser, wird das Wolfskind selbst zum Wolf.",
     "Chooses an enemy. If they die, Wolf Child becomes a werewolf."),

    ("Das Orakel", "The Oracle", "1.3",
     "Erfährt jede Nacht die Rolle eines Spielers.",
     "Each night learns the role of one chosen player."),

    ("Die Ewigen", "The Eternal Ones", "1.31",
     "Prüfen ob ein Spieler Solo-Siegbedingung hat. Gewinnen gemeinsam mit ihm.",
     "Each night: check if a player has a solo win condition. Win with them."),

    ("Spürhund", "Scent Hound", "1.4",
     "Wählt 3 Spieler. ✅ wenn ein Wolf/Solo dabei ist, sonst ❌.",
     "Picks 3 players. ✅ if a wolf or solo role is among them, else ❌."),

    ("Schutzengel", "Guardian Angel", "0.4",
     "Schützt jede Nacht einen Spieler vor dem Werwolfangriff.",
     "Each night protects one player from the werewolf attack."),

    ("Werwolf", "Werewolf", "0.6",
     "Tötet jede Nacht gemeinsam mit dem Rudel ein Opfer.",
     "Each night kills a victim together with the pack."),

    ("Rachsüchtiger Wolf", "Lone Wolf", "0.65",
     "Wacht jede dritte Nacht auf und kann einen anderen Wolf reißen. Solosieg.",
     "Wakes every third night and may kill another werewolf. Solo win."),

    ("König Lykaon", "King Lycaon", "0.7",
     "Verwandelt in Nacht 1 gemeinsam mit einem Verbündeten einen Dorfbewohner zum Wolf.",
     "Night 1: together with an ally turns one villager into a werewolf."),

    ("Der Weise", "The Elder", "X",
     "Überlebt den ersten Wolfsangriff. Bei Lynchung: Dorf verliert 1–3 Runden Fähigkeiten.",
     "Survives the first wolf attack. If lynched: village loses abilities for 1–3 rounds."),

    ("Verdammniswächter", "Doom Warden", "0.68",
     "Jede Nacht: wählt ob das Nachtopfer oder ein zufälliger anderer Spieler stirbt.",
     "Each night: chooses if the night victim or a random other player dies instead."),

    ("Lehrling", "Apprentice", "0.35",
     "Wählt einen Mentor. Übernimmt dessen Rolle nach seinem Tod.",
     "Chooses a mentor. Takes over their role when they die."),

    ("Busfahrer", "Bus Driver", "X",
     "Wenn gelyncht: sterben beide direkten Nachbarn mit.",
     "If lynched: both living neighbors die with them."),

    ("Korrupter Richter", "Corrupt Judge", "0.5",
     "Markiert einen Spieler der automatisch mit +1 Stimme nominiert ist.",
     "Marks a player who is automatically nominated with +1 vote."),

    ("Märtyrerin", "Martyr", "2.8",
     "Kann sich selbst opfern um das Nachtopfer zu retten.",
     "Can sacrifice herself to save the night victim."),

    ("Dorfwache", "Village Guard", "X",
     "Stirbt nicht wenn er nachts von den Werwölfen angegriffen wird.",
     "Does not die when targeted by werewolves at night."),

    ("Pestbringerin", "Plague Bringer", "1.7",
     "Verbreitet jede Nacht eine tödliche Seuche die sich weiter ausbreitet.",
     "Each night spreads a lethal plague that keeps spreading."),

    ("Prophet des Untergangs", "Prophet of Doom", "2.2",
     "Markiert 3 Spieler. Wenn alle tot: tötet er jede Nacht alleine.",
     "Marks 3 players. When all are dead: kills each night and wins alone."),

    ("Spiegelwolf", "Mirror Wolf", "X",
     "Beim ersten Lynch stirbt der Nominator statt ihm.",
     "On the first lynch: the nominator dies instead of him."),

    ("Dämonischer Wolf", "Demonic Wolf", "X",
     "Verflucht Opfer sodass sie bei Untersuchung als Werwölfe erscheinen.",
     "Curses victims so they appear as werewolves when investigated."),

    ("Trugbilderwolf", "Decoy Wolf", "X",
     "Zeigt dem Orakel eine zufällige Nicht-Wolf-Rolle anstatt Wolf.",
     "Shows the Oracle a random non-wolf role instead of wolf."),

    ("Schattenhund", "Shadow Hound", "0.2",
     "Einmalig: blockiert alle Dorf-Fähigkeiten für eine Nacht.",
     "Once per game: blocks all village abilities for one night."),

    ("Besessener Wolf", "Possessed Wolf", "X",
     "Reißt beim Tod (ab 5 Spielern) einen weiteren Spieler mit in den Tod.",
     "Upon death (≥5 players): drags another player to their death."),

    ("Fenrir", "Fenrir", "X",
     "Wird mit jeder überlebten Nacht stärker. Ab Stufe 3: überlebt einmalig jeden Tod.",
     "Grows stronger each survived night. From stage 3: survives any death once."),

    ("Kutscher", "Coachman", "X",
     "Ab 10 Toten: belebt 3 Tote wieder. Einer davon wird Wolf.",
     "With 10+ dead: revives 3 players. One of them becomes a werewolf."),

    ("Seelentauscher", "Soul Swapper", "1.9",
     "Einmalig: tauscht die Rollen zweier Spieler (lebend oder tot).",
     "Once per game: swaps the roles of two players, alive or dead."),

    ("Blutpriester", "Blood Priest", "2.0",
     "Opfert einen Spieler und deckt 0–3 Werwölfe auf.",
     "Sacrifices a player and reveals 0–3 werewolves."),

    ("Traumdeuter", "Dreamer", "1.5",
     "Erhält nächtliche Visionen über Rollen oder Zustände von Spielern.",
     "Receives nightly visions about the roles or states of players."),

    ("Henker", "Executioner", "1.85",
     "Nach 3 Lynchungen: markiert nachts ein Ziel das nach Lynchung zusätzlich stirbt.",
     "After 3 lynchings: marks a target each night who dies after the next lynch."),

    ("Feuerteufel", "Pyromaniac", "1.8",
     "Wählt ein Ziel. Stirbt es, verbrennen auch die benachbarten Spieler.",
     "Targets a player. When they die, their neighbors burn too."),

    ("Voodoo-Priester", "Voodoo Priest", "2.1",
     "Gibt einem Spieler eine Puppe. Stirbt er: stirbt stattdessen der Puppenträger.",
     "Gives a player a voodoo doll. If he would die, the doll holder dies instead."),

    ("Blutwolf", "Blood Wolf", "X",
     "Stimme zählt +1 für jeden direkt benachbarten toten Spieler.",
     "Their vote counts +1 for each directly adjacent dead neighbor."),

    ("Albtraumwolf", "Nightmare Wolf", "0.62",
     "Blockiert jede Nacht die Fähigkeit eines Dorfbewohners.",
     "Each night blocks the ability of one villager."),

    ("Cerberus", "Cerberus", "X",
     "Baut bis zu 3 Köpfe auf. Bei 3 Köpfen: kann eine Lynchung abwehren.",
     "Builds up to 3 heads. At 3 heads: can block a lynch."),

    ("Ritter", "Knight", "X",
     "Stirbt er nachts durch Wölfe: tötet er dabei den nächstliegenden Wolf.",
     "If killed by wolves at night: kills the nearest werewolf in return."),

    ("Rotkäppchen", "Little Red Riding Hood", "1.76",
     "Sucht jede Nacht Zuflucht. Gewährt Apfel-Buff und Todeskette mit dem Gastgeber.",
     "Seeks refuge nightly. Grants an apple buff and shares a death chain with the host."),

    ("Selbstmörder", "Death Seeker", "X",
     "Gewinnt wenn er bei mindestens 5 Toten am Tag gelyncht wird.",
     "Wins if lynched during the day when at least 5 players are already dead."),

    ("Kopfgeldjäger", "Bounty Hunter", "0.8",
     "Nach einem gelynchten Wolf: sieht 3 Spieler, einer davon ist ein Wolf.",
     "After a wolf is lynched: sees 3 players, one of whom is a werewolf."),

    ("König", "King", "1.2",
     "Wenn mehr Tote als Lebende: erfährt die Identität eines lebenden Dorfbewohners.",
     "When more players are dead than alive: learns a living villager's identity."),

    ("Dr. Victor Frankenstein", "Dr. Victor Frankenstein", "1.0",
     "Belebt einmalig einen toten Spieler und gibt ihm eine neue Rolle.",
     "Once revives a dead player who receives a brand new role."),

    ("Nekromant", "Necromancer", "0.75",
     "Kann Schutzschilde errichten oder Kills umlenken. Gewinnt allein durch Wolf-Benennung.",
     "Can raise shields or redirect kills. Wins alone by naming a living werewolf."),

    ("Kartenschlucker", "The Collector", "1.05",
     "Erhält Stapel wenn Tote Karten tauschen. Bei 10 Stapeln: sofortiger Sieg.",
     "Gains stacks when dead players exchange cards. At 10 stacks: wins immediately."),

    ("Hades", "Hades", "3.0",
     "Sammelt Lebenslichter von Toten. Kauft Fähigkeiten. Bei 10 Lichtern: Sieg.",
     "Collects life lights from the dead. Buys abilities. Wins at 10 lights."),

    ("Siegreicher Wolf", "Victorious Wolf", "X",
     "Solange er lebt: zählt er für die Wolfssiegbedingung wie zwei Werwölfe.",
     "While alive: counts as two werewolves toward the wolf win condition."),

    ("Seuchenwolf", "Blight Wolf", "X",
     "Nach seinem Tod: nächster Wolfsangriff durchdringt alle Schutzeffekte.",
     "After their death: the next wolf attack pierces all protection effects."),

    ("Schicksalswolf", "Fate Wolf", "0.705",
     "Wählt 3 Spieler. Für jeden unter den ersten 3 Toten: Extra-Kill in Nacht 4.",
     "Picks 3 players. For each among the first 3 dead: gains a kill in night 4."),

    ("Schattenwanderer", "Shadowwalker", "0.708",
     "Knüpft Todeskette mit einem Spieler. Stirbt einer, stirbt der andere stattdessen.",
     "Creates a death chain with one player. If one dies, the other dies instead."),

    ("Giftwolf", "Poison Wolf", "0.709",
     "Zweimal pro Spiel: vergiftet ein Ziel. Ziel stirbt zwei Tage später.",
     "Twice per game: poisons a target. The target dies two days later."),

    ("Rudelvater", "Packfather", "X",
     "Überlebt ersten nicht-Wolf/Lynch-Tod. Bei Lynchung: Wölfe erhalten Zusatzangriff.",
     "Survives the first non-wolf/lynch death. If lynched: wolves gain a second attack."),

    ("Schwarze Witwe", "Black Widow", "0.7095",
     "Wählt jede Nacht einen Spieler. Findet sie Liebende/Rivalen: sterben beide.",
     "Each night picks a player. If a lover or rival is found, both die the next day."),

    ("Doktor", "Doctor", "1.315",
     "Nimmt Blutproben von 2 Spielern. Erfährt ob sie demselben Team angehören.",
     "Takes blood samples from 2 players. Learns if they belong to the same team."),

    ("Fährtenleser", "Tracker", "1.316",
     "Einmalig: erfährt ob der nächstliegende Wolf links oder rechts sitzt.",
     "Once per game: learns if the nearest wolf sits to the left or right."),

    ("Waldläufer", "Ranger", "1.317",
     "Erfährt jede Nacht wie viele lebende Werwölfe im Spiel sind.",
     "Each night learns the exact number of living werewolves in the game."),

    ("Schutzgeist", "Guardian Spirit", "1.318",
     "In der Nacht nach ihrem Tod: schützt einen Spieler mit einem Schild.",
     "On the night after death: shields a chosen player from the next attack."),

    ("Dorfchronistin", "Village Chronicler", "0.12",
     "Erfährt zu Spielbeginn wie viele Solo-Rollen im Spiel sind.",
     "Learns at the start of the game how many solo roles are in play."),

    ("Wächter am Tor", "Gatewarden", "X",
     "Solange er lebt: werden neu entstehende Werwölfe stattdessen zu Dorfbewohnern.",
     "While alive: newly created werewolves become villagers instead."),

    ("Zeitwächter", "Time Warden", "2.8",
     "Einmalig: friert eine Nacht ein. Alle Aktionen werden abgebrochen.",
     "Once per game: freezes a night — all night actions are canceled."),

    ("Amalia", "Amalia", "1.3195",
     "Opfert sich um eine öffentliche Ja/Nein-Frage zu stellen (bei 2+ Wölfen).",
     "Sacrifices herself to publicly ask a yes/no question (while 2+ wolves live)."),

    ("Kriegerin des Lichts", "Warrior of Light", "1.3197",
     "Einmalig: greift nachts an. Kein Wolf getroffen = sie stirbt selbst.",
     "Once per game: attacks at night. If not a wolf: she dies instead."),

    ("Detektiv", "Detective", "X",
     "Nach dem Tod eines Wolfes: öffentlicher Hinweis auf einen anderen Wolf.",
     "After a wolf dies: a public clue about another wolf is revealed."),

    ("Dorfschmied", "Village Blacksmith", "0.55",
     "Schmiedet 5 Nächte. Nacht 6: gibt Waffe weiter die einen Wolf tötet.",
     "Forges for 5 nights. Night 6: gives a weapon that kills a random wolf."),

    ("Manipulator", "Manipulator", "X",
     "Gewinnt wenn er Final 3 erreicht ohne je nominiert zu sein. Nominierung = Tod.",
     "Wins if reaching final 3 without ever being nominated. Nomination means death."),

    ("Doppelspion", "Double Agent", "0.605",
     "Wacht mit Werwölfen auf. Gewinnt allein wenn alle Werwölfe tot sind.",
     "Wakes with werewolves. Wins alone when all werewolves are dead."),

    ("Grabräuber", "Grave Robber", "1.3199",
     "Einmalig: stiehlt die Fähigkeit eines toten Spielers. Solosieg.",
     "Once per game: steals a dead player's ability. Solo win condition."),

    ("Parasit", "Parasite", "1.3198",
     "Heftet sich jede Nacht an einen Spieler. Stirbt nur wenn sein Wirt stirbt.",
     "Attaches to a living player each night. Only dies when the host dies."),

    ("Todesprediger", "Death Prophet", "1.31995",
     "Prophezeit seinen eigenen Todeszeitpunkt. Liegt er richtig: Solosieg.",
     "Predicts the exact moment of their own death. If correct: wins alone."),

    ("Dorfbewohner", "Villager", "X",
     "Hat keine aktive Nachtfähigkeit.",
     "Has no active night ability."),
]

# ─── HILFSFUNKTIONEN ──────────────────────────────────────────────────────────

def set_cell_bg(cell, hex_color):
    """Setzt Hintergrundfarbe einer Tabellenzelle."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def add_header_row(table, texts, bg='1C1C1C', fg='C9A84C'):
    """Fügt eine Kopfzeile zur Tabelle hinzu."""
    row = table.rows[0]
    for i, text in enumerate(texts):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xC9, 0xA8, 0x4C)
        run.font.name = 'Calibri'

def add_data_row(table, cells_data, row_idx):
    """Fügt eine Datenzeile ein."""
    row = table.add_row()
    bg = '2A2A2A' if row_idx % 2 == 0 else '222222'
    for i, (text, bold) in enumerate(cells_data):
        cell = row.cells[i]
        set_cell_bg(cell, bg)
        p = cell.paragraphs[0]
        if i == 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.font.size = Pt(8.5)
        run.font.name = 'Calibri'
        run.bold = bold
        if i == 0:
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif i == 1:
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        elif i == 2:
            if text == 'X':
                run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            else:
                run.font.color.rgb = RGBColor(0xC9, 0xA8, 0x4C)
        else:
            run.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)

# ─── DOKUMENT ERSTELLEN ───────────────────────────────────────────────────────

doc = Document()

# Seitenränder
section = doc.sections[0]
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)

# Titel
title = doc.add_heading('Grimmhain – Werewolf Reckoning', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = RGBColor(0xC9, 0xA8, 0x4C)

subtitle = doc.add_paragraph('Alle Rollen · Night Order Tier · Kurzfähigkeit')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_run = subtitle.runs[0]
sub_run.font.size = Pt(10)
sub_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

doc.add_paragraph('')

# Hinweis
note = doc.add_paragraph('Tier X = keine Nachtphase-Aktion  |  Rollen nach Tier aufsteigend sortiert')
note.alignment = WD_ALIGN_PARAGRAPH.CENTER
note_run = note.runs[0]
note_run.font.size = Pt(8)
note_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
note_run.italic = True

doc.add_paragraph('')

# Tabelle: 5 Spalten
# Rollenname DE | Rollenname EN | Tier | Beschreibung DE | Beschreibung EN
headers = ['Rollenname (DE)', 'Role Name (EN)', 'Tier', 'Fähigkeit (DE)', 'Ability (EN)']
table = doc.add_table(rows=1, cols=5)
table.style = 'Table Grid'

# Spaltenbreiten
col_widths = [Cm(3.2), Cm(3.5), Cm(1.2), Cm(7.5), Cm(7.5)]
for i, width in enumerate(col_widths):
    for cell in table.columns[i].cells:
        cell.width = width

# Header
add_header_row(table, headers)

# Daten – sortiert nach Tier (X ganz unten)
def sort_key(role):
    tier = role[2]
    if tier == 'X':
        return (1, 0)
    return (0, float(tier))

sorted_roles = sorted(ROLES, key=sort_key)

for idx, (de_name, en_name, tier, de_desc, en_desc) in enumerate(sorted_roles):
    cells_data = [
        (de_name, True),
        (en_name, False),
        (tier, True),
        (de_desc, False),
        (en_desc, False),
    ]
    add_data_row(table, cells_data, idx)

# Ausgabepfad
output_path = os.path.join(os.path.expanduser('~'), 'Downloads', 'Grimmhain_Rollen_Übersicht.docx')
doc.save(output_path)
print(f"Gespeichert: {output_path}")
print(f"Rollen gesamt: {len(ROLES)}")
